"""Excel-отчёт в рабочем формате: Calibri 10, без заливок и декора.

Лист "Общие итоги": таблица Показатель / Значение / Комментарий (ширины 46 / 12 / 72).
На листах с итоговой строкой она стоит над шапкой, на шапке включён фильтр.
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src import config
from src.plots import SHORT, fmt_pct, fmt_rub

FONT = Font(name="Calibri", size=10)
BOLD = Font(name="Calibri", size=10, bold=True)
RUB = "#,##0"
PCT = "0.0%"


def _p(p: float) -> str:
    """p-value в русской записи."""
    return "менее 0,001" if p < 0.001 else f"{p:.3f}".replace(".", ",")


def _write_table(ws, df: pd.DataFrame, formats: dict[str, str], widths: list[int],
                 totals: dict[str, object] | None = None) -> None:
    """Таблица с необязательной строкой итогов над шапкой и фильтром на шапке."""
    header_row = 2 if totals else 1
    if totals:
        for j, col in enumerate(df.columns, start=1):
            if col in totals:
                c = ws.cell(row=1, column=j, value=totals[col])
                c.font = BOLD
                c.number_format = formats.get(col, "General")
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=header_row, column=j, value=col).font = BOLD
    for i, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for j, (col, v) in enumerate(zip(df.columns, row), start=1):
            c = ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
            c.font = FONT
            c.number_format = formats.get(col, "General")
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str))
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    last = get_column_letter(len(df.columns))
    ws.auto_filter.ref = f"A{header_row}:{last}{header_row + len(df)}"
    ws.freeze_panes = f"A{header_row + 1}"


def build_report(pnl: pd.DataFrame, effects: pd.DataFrame, plan: pd.DataFrame,
                 pretrends: pd.DataFrame, wape_naive: float, wape_did: float,
                 data_info: dict) -> str:
    wb = Workbook()

    # Лист 1: Общие итоги
    ws = wb.active
    ws.title = "Общие итоги"
    rows = [
        ("Период данных", data_info["period"], f"{config.N_WEEKS} недель, {config.N_STORES} магазинов, синтетические данные"),
        ("Чеков в выгрузке", data_info["receipts"], f"строк чеков: {fmt_rub(data_info['lines'])}"),
    ]
    for r in pnl.itertuples():
        rows.append((
            f"Результат акции: {r.name}", round(r.incremental_pnl),
            f"{r.status}; маржа за волну {fmt_rub(r.incr_margin)} руб, 95% интервал "
            f"от {fmt_rub(r.incr_margin_ci_low)} до {fmt_rub(r.incr_margin_ci_high)} руб, "
            f"p-value {_p(r.p_value)}; фиксированные затраты {fmt_rub(r.fixed_cost)} руб"))
    total = pnl.incremental_pnl.round().sum()
    loss = pnl.loc[pnl.status == "убыточная", "incremental_pnl"].round().sum()
    rows += [
        ("Итого результат всех акций первой волны, руб", round(total),
         "сумма инкрементальных результатов после фиксированных затрат"),
        ("Потери на убыточных акциях, руб", round(loss),
         "деньги, которые можно перевести в работающие механики"),
        ("Ошибка плана второй волны: оценка до/после", wape_naive,
         "WAPE по продажам участников, база продаж одинаковая для обоих подходов"),
        ("Ошибка плана второй волны: оценка DiD", wape_did,
         "WAPE по продажам участников; эффект первой волны очищен от сезонности"),
        ("Параллельные тренды до акций", f"{int(pretrends.parallel_trends_ok.sum())} из {len(pretrends)}",
         "совместный тест недель до старта, p-value больше 0,05 у всех акций"
         if pretrends.parallel_trends_ok.all() else "есть акции с расхождением трендов до старта"),
    ]
    df = pd.DataFrame(rows, columns=["Показатель", "Значение", "Комментарий"])
    _write_table(ws, df, {}, [46, 12, 72])
    for i, v in enumerate(df["Значение"], start=2):
        c = ws.cell(row=i, column=2)
        if isinstance(v, float) and abs(v) < 1:
            c.number_format = PCT
        elif isinstance(v, (int, float)):
            c.number_format = RUB

    # Лист 2: Акции, итоги над шапкой
    ws = wb.create_sheet("Акции")
    t = pd.DataFrame({
        "Акция": pnl.name,
        "Магазинов": pnl.stores,
        "Недель": pnl.weeks,
        "Доп. маржа, руб": pnl.incr_margin.round(),
        "95% от, руб": pnl.incr_margin_ci_low.round(),
        "95% до, руб": pnl.incr_margin_ci_high.round(),
        "p-value": pnl.p_value.round(4),
        "Фикс. затраты, руб": pnl.fixed_cost,
        "Результат, руб": pnl.incremental_pnl.round(),
        "Скидки и подарки, руб": pnl.discount_and_gifts.round(),
        "Результат на 100 руб скидки": pnl.pnl_per_100_rub_given.round(1),
        "Статус": pnl.status,
    })
    totals = {"Акция": "Итого", "Доп. маржа, руб": t["Доп. маржа, руб"].sum(),
              "Фикс. затраты, руб": t["Фикс. затраты, руб"].sum(),
              "Результат, руб": t["Результат, руб"].sum(),
              "Скидки и подарки, руб": t["Скидки и подарки, руб"].sum()}
    fm = {c: RUB for c in t.columns if "руб" in c and "100" not in c}
    fm["Результат на 100 руб скидки"] = "0.0"
    fm["p-value"] = "0.0000"
    _write_table(ws, t, fm, [34, 11, 9, 15, 13, 13, 9, 16, 15, 19, 16, 20], totals)

    # Лист 3: Эффекты DiD
    ws = wb.create_sheet("Эффекты DiD")
    e = pd.DataFrame({
        "Акция": effects.promo_id.map(SHORT),
        "Показатель": effects.label,
        "Эффект на магазин в неделю": effects.coef.round(1),
        "Эффект, %": effects.effect_pct,
        "95% от": effects.ci_low.round(1),
        "95% до": effects.ci_high.round(1),
        "p-value": effects.p_value.round(4),
    })
    _write_table(ws, e, {"Эффект, %": PCT, "p-value": "0.0000", "Эффект на магазин в неделю": "#,##0.0",
                         "95% от": "#,##0.0", "95% до": "#,##0.0"}, [22, 40, 16, 11, 12, 12, 9])

    # Лист 4: Планирование, итоги над шапкой
    ws = wb.create_sheet("Планирование")
    p = pd.DataFrame({
        "Акция, волна 2": plan.promo_id.map(SHORT),
        "Эффект до/после": plan.uplift_naive,
        "Эффект DiD": plan.uplift_did,
        "База, шт": plan.baseline_units.round(),
        "План до/после, шт": plan.plan_naive.round(),
        "План DiD, шт": plan.plan_did.round(),
        "Факт, шт": plan.actual_units,
        "Ошибка до/после": plan.ape_naive,
        "Ошибка DiD": plan.ape_did,
    })
    totals = {"Акция, волна 2": "Итого (WAPE)", "База, шт": p["База, шт"].sum(),
              "План до/после, шт": p["План до/после, шт"].sum(), "План DiD, шт": p["План DiD, шт"].sum(),
              "Факт, шт": p["Факт, шт"].sum(), "Ошибка до/после": wape_naive, "Ошибка DiD": wape_did}
    fm = {c: PCT for c in ("Эффект до/после", "Эффект DiD", "Ошибка до/после", "Ошибка DiD")}
    fm.update({c: RUB for c in ("База, шт", "План до/после, шт", "План DiD, шт", "Факт, шт")})
    _write_table(ws, p, fm, [22, 15, 12, 11, 17, 13, 11, 15, 12], totals)

    config.REPORTS.mkdir(parents=True, exist_ok=True)
    path = config.REPORTS / "promo_summary.xlsx"
    wb.save(path)
    return str(path)
